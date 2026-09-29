"""Export helpers: CSV (streaming) + XLSX (openpyxl).

PDF is an explicit boundary: no pure-python PDF writer is installed
(reportlab is absent from requirements and the venv), so this module ships
CSV/XLSX only and reports PDF as unsupported instead of adding a heavy dep.
"""
from __future__ import annotations

import csv
import io
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:  # pragma: no cover - import probe, openpyxl is a hard dep
    import openpyxl  # type: ignore
    from openpyxl.styles import Font  # type: ignore

    _HAS_OPENPYXL = True
except Exception:  # pragma: no cover
    _HAS_OPENPYXL = False

SUPPORTED_FORMATS = ("csv", "xlsx")

CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

PDF_SUPPORTED = False
PDF_REASON = (
    "PDF export is not shipped: no pure-python PDF writer is installed "
    "(reportlab is absent from requirements and the runtime). CSV and XLSX "
    "cover the tabular export contract; add a PDF writer as an explicit "
    "dependency before enabling this path."
)


def pdf_status() -> Dict[str, Any]:
    """Describe the PDF boundary for callers and docs."""
    return {"supported": PDF_SUPPORTED, "reason": PDF_REASON}


def normalize_rows(rows: Any) -> List[Dict[str, Any]]:
    """Coerce export input to a list of plain dicts (deterministic)."""
    if rows is None:
        return []
    if hasattr(rows, "to_dict"):  # DataFrame
        try:
            records = rows.to_dict("records")  # type: ignore[union-attr]
            return [{str(k): v for k, v in dict(r).items()} for r in records]
        except Exception:
            return []
    if isinstance(rows, dict):
        return [{str(k): v for k, v in rows.items()}]
    out: List[Dict[str, Any]] = []
    for row in list(rows):
        if isinstance(row, dict):
            out.append({str(k): v for k, v in row.items()})
        else:
            out.append({"value": row})
    return out


def resolve_columns(rows: List[Dict[str, Any]], columns: Any = None) -> List[str]:
    """Column order: explicit list wins, else first-seen union (deterministic)."""
    if columns:
        cols = [str(c) for c in list(columns) if str(c).strip()]
        if cols:
            return cols
    seen: List[str] = []
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.append(key)
    return seen


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def to_csv_bytes(rows: Any, columns: Any = None) -> bytes:
    """Render rows to UTF-8-SIG CSV bytes (Excel-friendly)."""
    records = normalize_rows(rows)
    cols = resolve_columns(records, columns)
    buf = io.StringIO(newline="")
    writer = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore", lineterminator="\r\n")
    if cols:
        writer.writeheader()
        for row in records:
            writer.writerow({c: _cell(row.get(c, "")) for c in cols})
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def iter_csv(rows: Any, columns: Any = None, chunk_size: int = 1000) -> Iterable[bytes]:
    """Stream CSV in chunks for large exports (header first, then row blocks)."""
    records = normalize_rows(rows)
    cols = resolve_columns(records, columns)
    header = io.StringIO(newline="")
    writer = csv.DictWriter(header, fieldnames=cols, extrasaction="ignore", lineterminator="\r\n")
    first = True
    if cols:
        writer.writeheader()
        text = header.getvalue()
        yield (("\ufeff" + text) if first else text).encode("utf-8")
        first = False
    else:
        yield "﻿".encode("utf-8")
        return
    try:
        n = max(1, int(chunk_size))
    except (TypeError, ValueError):
        n = 1000
    for i in range(0, len(records), n):
        block = io.StringIO(newline="")
        w = csv.DictWriter(block, fieldnames=cols, extrasaction="ignore", lineterminator="\r\n")
        for row in records[i:i + n]:
            w.writerow({c: _cell(row.get(c, "")) for c in cols})
        yield block.getvalue().encode("utf-8")


def to_xlsx_bytes(rows: Any, columns: Any = None, sheet_name: str = "data") -> bytes:
    """Render rows to XLSX bytes via openpyxl."""
    if not _HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required for XLSX export")
    records = normalize_rows(rows)
    cols = resolve_columns(records, columns)
    import openpyxl as _oxl

    wb = _oxl.Workbook()
    ws = wb.active
    ws.title = str(sheet_name or "data")[:31]
    if cols:
        for j, col in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=j, value=col)
            cell.font = Font(bold=True)
        for i, row in enumerate(records, start=2):
            for j, col in enumerate(cols, start=1):
                ws.cell(row=i, column=j, value=_cell(row.get(col, "")))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_rows(
    rows: Any,
    fmt: str = "csv",
    columns: Any = None,
    filename: Any = None,
) -> Tuple[bytes, str, str]:
    """Export rows to (content, media_type, filename). Raises ValueError on fmt."""
    kind = str(fmt or "").strip().lower()
    if kind not in SUPPORTED_FORMATS:
        raise ValueError("format must be one of: csv, xlsx (pdf is not supported: " + PDF_REASON + ")")
    if kind == "pdf":
        raise ValueError(PDF_REASON)  # pragma: no cover - unreachable via guard above
    records = normalize_rows(rows)
    cols = resolve_columns(records, columns)
    base = str(filename or "export").strip() or "export"
    base = "".join(c for c in base if c.isalnum() or c in ("-", "_", ".", " ")).strip() or "export"
    if kind == "csv":
        return to_csv_bytes(records, cols), CSV_MEDIA_TYPE, f"{base}.csv"
    return to_xlsx_bytes(records, cols), XLSX_MEDIA_TYPE, f"{base}.xlsx"
